"""Export one editable scene to native PPTX, vector SVG, and matching PNG. No model/API calls."""
import argparse
from collections import Counter
from copy import deepcopy
from html import escape
import json
import math
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

COLOR = re.compile(r'^#[0-9a-fA-F]{6}$')
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def color(value):
    return RGBColor.from_string(value.lstrip('#'))


def font_path(family, bold=False, italic=False):
    windows = Path('C:/Windows/Fonts')
    arial = 'arialbi.ttf' if bold and italic else 'arialbd.ttf' if bold else 'ariali.ttf' if italic else 'arial.ttf'
    for filename in ((('msyhbd.ttc' if bold else 'msyh.ttc'), 'simhei.ttf') if any(x in family.lower() for x in ['yahei', 'hei', 'cjk']) else (arial, 'arial.ttf', 'calibri.ttf')):
        if (windows / filename).exists():
            return windows / filename
    for filename in ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'):
        if Path(filename).exists():
            return Path(filename)
    return None


def wrap_text(node, family):
    text = str(node['text'])
    font_file = font_path(family, node.get('bold', False), node.get('italic', False))
    font = ImageFont.truetype(str(font_file), max(1, round(node.get('fontSize', 22)))) if font_file else ImageFont.load_default()
    available = node['w'] - 2 * node.get('padding', 1)
    if node.get('wrap', True) is False:
        return text.split('\n'), font
    lines = []
    for paragraph in text.split('\n'):
        current = ''
        # Preserve whitespace boundaries for Latin scripts, allow CJK breaks.
        for token in re.findall(r'[\u2e80-\uffff]|[^\S\n]+|[^\s\u2e80-\uffff]+', paragraph):
            candidate = current + token
            if current and font.getlength(candidate) > available:
                lines.append(current.rstrip())
                current = token.lstrip()
            else:
                current = candidate
        lines.append(current.rstrip())
    return lines, font


def prepare(scene):
    scene = deepcopy(scene)
    errors, warnings, ids = [], [], set()
    width, height = scene['canvas']['width'], scene['canvas']['height']
    if not 0 < width <= 10000 or not 0 < height <= 10000:
        raise ValueError('Canvas dimensions must be positive and <= 10000.')
    family = scene['canvas'].get('fontFamily', 'Arial')
    for collection in ('nodes', 'edges'):
        for item in scene.get(collection, []):
            ident = item.get('id')
            if not ident or ident in ids:
                errors.append(f'Missing or duplicate id: {ident}')
            ids.add(ident)
            for key in ('fill', 'stroke', 'color'):
                value = item.get(key)
                if value is not None and value != 'none' and not COLOR.fullmatch(value):
                    errors.append(f'{ident}: {key} must be #rrggbb or none')
    nodes = {node['id']: node for node in scene['nodes']}
    for node in scene['nodes']:
        ident, kind = node['id'], node['type']
        if kind not in {'rect', 'roundrect', 'ellipse', 'polygon', 'text', 'path'}:
            errors.append(f'{ident}: unsupported node type {kind}; do not replace it with a raster snapshot')
            continue
        if kind == 'path':
            coordinates = []
            for command in node.get('commands', []):
                expected = {'M': 2, 'L': 2, 'C': 6, 'Q': 4, 'Z': 0}.get(command[0])
                if expected is None or len(command)-1 != expected:
                    errors.append(f'{ident}: invalid path command')
                    continue
                coordinates.extend(zip(command[1::2], command[2::2]))
            if not coordinates:
                errors.append(f'{ident}: path has no coordinates')
                continue
            xs, ys = zip(*coordinates)
            node.update(x=min(xs), y=min(ys), w=max(.01,max(xs)-min(xs)), h=max(.01,max(ys)-min(ys)))
        if kind == 'polygon':
            points = node['points']
            if len(points) < 3:
                errors.append(f'{ident}: polygon needs at least 3 vertices')
            xs, ys = [p[0] for p in points], [p[1] for p in points]
            node.update(x=min(xs), y=min(ys), w=max(xs) - min(xs), h=max(ys) - min(ys))
        if min(node.get('w', 0), node.get('h', 0)) <= 0:
            errors.append(f'{ident}: nonpositive dimensions')
        if node.get('x', 0) < 0 or node.get('y', 0) < 0 or node.get('x', 0) + node.get('w', 0) > width + .1 or node.get('y', 0) + node.get('h', 0) > height + .1:
            errors.append(f'{ident}: outside canvas')
        if kind == 'text':
            lines, font = wrap_text(node, family)
            node['_lines'] = lines
            fs = node.get('fontSize', 22)
            leading = node.get('lineHeight', 1.22) * fs
            padding = node.get('padding', 1)
            if len(lines) * leading + 2 * padding > node['h'] + .1:
                errors.append(f'{ident}: text height exceeds box ({len(lines)} lines)')
            if any(font.getlength(line) > node['w'] - 2 * padding + .1 for line in lines):
                errors.append(f'{ident}: text width exceeds box')
            if fs < scene['canvas'].get('minFontSize', 16):
                warnings.append(f'{ident}: font below configured publication minimum')
        if node.get('gradient'):
            gradient = node['gradient']
            if gradient.get('angle',90) not in (0,90) or len(gradient.get('colors',[])) < 2 or any(not COLOR.fullmatch(c) for c in gradient.get('colors',[])):
                errors.append(f'{ident}: gradient needs 2+ hex colors and angle 0 or 90')
    for edge in scene.get('edges', []):
        if 'points' not in edge:
            for key in ('from', 'to'):
                if edge.get(key) not in nodes or nodes.get(edge.get(key), {}).get('type') == 'text':
                    errors.append(f'{edge["id"]}: {key} must reference a non-text node')
            if edge.get('from') in nodes and edge.get('to') in nodes:
                edge['points'] = [anchor(nodes[edge['from']], edge.get('fromSide', 'right')), anchor(nodes[edge['to']], edge.get('toSide', 'left'))]
        if len(edge.get('points', [])) < 2:
            errors.append(f'{edge["id"]}: edge needs >=2 points')
        if any(not (0 <= x <= width and 0 <= y <= height) for x, y in edge.get('points', [])):
            errors.append(f'{edge["id"]}: edge outside canvas')
    if errors:
        raise ValueError('\n'.join(errors))
    return scene, warnings


def anchor(node, side):
    x, y, w, h = [node[key] for key in ('x', 'y', 'w', 'h')]
    return {'left': [x, y + h/2], 'right': [x+w, y+h/2], 'top': [x+w/2, y], 'bottom': [x+w/2, y+h]}[side]


def svg_node(node, family):
    kind, ident = node['type'], escape(node['id'], quote=True)
    x, y, w, h = [node[key] for key in ('x', 'y', 'w', 'h')]
    if kind == 'text':
        fs, padding = node.get('fontSize', 22), node.get('padding', 1)
        leading = fs * node.get('lineHeight', 1.22)
        align = node.get('align', 'left')
        tx = x + (w/2 if align == 'center' else w-padding if align == 'right' else padding)
        anchor_value = {'left': 'start', 'center': 'middle', 'right': 'end'}[align]
        text_height = len(node['_lines']) * leading
        top = y + (h-text_height)/2 if node.get('valign', 'middle') == 'middle' else y+padding
        lines = ''.join(f'<tspan x="{tx}" y="{top + fs*.93 + i*leading}">{escape(line)}</tspan>' for i, line in enumerate(node['_lines']))
        return f'<text id="{ident}" font-family="{escape(family, quote=True)}" font-size="{fs}" font-weight="{700 if node.get("bold") else 400}" font-style="{"italic" if node.get("italic") else "normal"}" text-anchor="{anchor_value}" fill="{node.get("color", "#172337")}">{lines}</text>'
    fill = f'url(#gradient-{ident})' if node.get('gradient') else node.get('fill','none')
    style = f'fill="{fill}" stroke="{node.get("stroke", "none")}" stroke-width="{node.get("strokeWidth", 1.5)}"'
    if node.get('dash'):
        style += ' stroke-dasharray="7 5"'
    if kind in {'rect', 'roundrect'}:
        return f'<rect id="{ident}" x="{x}" y="{y}" width="{w}" height="{h}" rx="{node.get("radius", 12) if kind=="roundrect" else 0}" {style}/>'
    if kind == 'ellipse':
        return f'<ellipse id="{ident}" cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" {style}/>'
    if kind == 'path':
        data = ' '.join(command[0]+' '+ ' '.join(str(round(v,4)) for v in command[1:]) for command in node['commands'])
        return f'<path id="{ident}" d="{data}" {style} stroke-linejoin="round" stroke-linecap="round"/>'
    points = ' '.join(f'{a},{b}' for a, b in node['points'])
    return f'<polygon id="{ident}" points="{points}" {style}/>'


def svg_edge(edge):
    points = ' '.join(f'{x},{y}' for x, y in edge['points'])
    stroke = edge.get('stroke', '#4B5B70')
    marker = f' marker-end="url(#arrow-{stroke[1:]})"' if edge.get('arrow', True) else ''
    dash = ' stroke-dasharray="7 5"' if edge.get('dash') else ''
    return f'<polyline id="{escape(edge["id"], quote=True)}" points="{points}" fill="none" stroke="{stroke}" stroke-width="{edge.get("strokeWidth", 2)}" stroke-linejoin="round"{marker}{dash}/>'


def elements(scene):
    # Stable z-order shared by SVG and PPTX; panels behind wires, labels in front.
    return sorted([(n.get('z', 0 if n['type'] != 'text' else 30), 'node', n) for n in scene['nodes']] + [(e.get('z', 20), 'edge', e) for e in scene.get('edges', [])], key=lambda item: item[0])


def build_svg(scene):
    canvas = scene['canvas']
    width, height = canvas['width'], canvas['height']
    strokes = {edge.get('stroke', '#4B5B70') for edge in scene.get('edges', [])}
    definitions = ''.join(f'<marker id="arrow-{stroke[1:]}" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="{stroke}"/></marker>' for stroke in strokes)
    for node in scene['nodes']:
        if node.get('gradient'):
            gradient=node['gradient']; colors=gradient['colors']
            axis='x2="0%" y2="100%"' if gradient.get('angle',90)==90 else 'x2="100%" y2="0%"'
            stops=''.join(f'<stop offset="{100*i/(len(colors)-1)}%" stop-color="{value}"/>' for i,value in enumerate(colors))
            definitions+=f'<linearGradient id="gradient-{escape(node["id"],quote=True)}" x1="0%" y1="0%" {axis}>{stops}</linearGradient>'
    body = ''.join(svg_node(item, canvas.get('fontFamily', 'Arial')) if kind == 'node' else svg_edge(item) for _, kind, item in elements(scene))
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}" role="img"><title>{escape(scene.get("title", "Research architecture"))}</title><defs>{definitions}</defs><rect width="{width}" height="{height}" fill="{canvas.get("background", "#FFFFFF")}"/>{body}</svg>'


def style_shape(shape, node, scale):
    # Explicitly suppress the PowerPoint theme's default shadow/effect preset.
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    fill, stroke = node.get('fill', 'none'), node.get('stroke', 'none')
    if hasattr(shape, 'fill'):
        if fill == 'none':
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = color(fill)
        if node.get('gradient'):
            fill_element=shape._element.spPr.find('{'+NS['a']+'}solidFill')
            fill_index=shape._element.spPr.index(fill_element) if fill_element is not None else 2
            if fill_element is not None: shape._element.spPr.remove(fill_element)
            gradient=OxmlElement('a:gradFill'); gradient.set('rotWithShape','1')
            stops=OxmlElement('a:gsLst'); colors=node['gradient']['colors']
            for index,value in enumerate(colors):
                stop=OxmlElement('a:gs'); stop.set('pos',str(round(index*100000/(len(colors)-1))))
                rgb=OxmlElement('a:srgbClr'); rgb.set('val',value[1:]); stop.append(rgb); stops.append(stop)
            gradient.append(stops)
            linear=OxmlElement('a:lin'); linear.set('ang',str(node['gradient'].get('angle',90)*60000)); linear.set('scaled','1'); gradient.append(linear)
            shape._element.spPr.insert(fill_index,gradient)
    if stroke == 'none':
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(node.get('strokeWidth', 1.5) * scale * 72)
    if node.get('dash'):
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def build_pptx(scene, destination):
    canvas = scene['canvas']
    scale = 13.333333 / canvas['width']
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(canvas['width'] * scale), Inches(canvas['height'] * scale)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(canvas.get('background', '#FFFFFF'))
    created, groups = {}, {}
    for _, kind, node in elements(scene):
        if kind == 'edge':
            points = node['points']
            for index, (start, end) in enumerate(zip(points, points[1:])):
                connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(start[0]*scale), Inches(start[1]*scale), Inches(end[0]*scale), Inches(end[1]*scale))
                connector.name = f'{node["id"]}:{index}'
                style_shape(connector, {**node, 'fill':'none', 'stroke': node.get('stroke', '#4B5B70')}, scale)
                if node.get('arrow', True) and index == len(points)-2:
                    head = OxmlElement('a:tailEnd')
                    head.set('type', 'triangle')
                    head.set('w', 'med')
                    head.set('len', 'med')
                    connector._element.spPr.get_or_add_ln().append(head)
                # With computed endpoints, preserve native endpoint attachment.
                sides = {'top': 0, 'left': 1, 'bottom': 2, 'right': 3}
                if index == 0 and node.get('from') in created:
                    connector.begin_connect(created[node['from']], sides[node.get('fromSide', 'right')])
                if index == len(points)-2 and node.get('to') in created:
                    connector.end_connect(created[node['to']], sides[node.get('toSide', 'left')])
            continue
        x, y, w, h = [node[key] for key in ('x','y','w','h')]
        if node['type'] == 'text':
            shape = slide.shapes.add_textbox(Inches(x*scale), Inches(y*scale), Inches(w*scale), Inches(h*scale))
            frame = shape.text_frame
            frame.clear()
            frame.word_wrap = False
            padding = Inches(node.get('padding', 1)*scale)
            frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = padding
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE if node.get('valign', 'middle') == 'middle' else MSO_ANCHOR.TOP
            fs = node.get('fontSize', 22)
            for index, line in enumerate(node['_lines']):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.alignment = {'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[node.get('align', 'left')]
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = Pt(fs * node.get('lineHeight', 1.22) * scale * 72)
                run = paragraph.add_run()
                run.text = line
                run.font.name = canvas.get('fontFamily','Arial')
                east_asian=OxmlElement('a:ea'); east_asian.set('typeface',canvas.get('fontFamily','Arial')); run._r.get_or_add_rPr().append(east_asian)
                run.font.size = Pt(fs*scale*72)
                run.font.bold = node.get('bold',False)
                run.font.italic = node.get('italic',False)
                run.font.color.rgb = color(node.get('color','#172337'))
        elif node['type'] == 'path':
            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x*scale), Inches(y*scale), Inches(w*scale), Inches(h*scale))
            preset=shape._element.spPr.find('{'+NS['a']+'}prstGeom')
            if preset is not None: shape._element.spPr.remove(preset)
            geometry=OxmlElement('a:custGeom')
            for tag in ('avLst','gdLst','ahLst','cxnLst'): geometry.append(OxmlElement('a:'+tag))
            rectangle=OxmlElement('a:rect')
            for key,value in {'l':'0','t':'0','r':'r','b':'b'}.items(): rectangle.set(key,value)
            geometry.append(rectangle)
            paths=OxmlElement('a:pathLst'); vector=OxmlElement('a:path')
            vector.set('w',str(max(1,round(w*1000)))); vector.set('h',str(max(1,round(h*1000))))
            for command in node['commands']:
                tag={'M':'moveTo','L':'lnTo','C':'cubicBezTo','Q':'quadBezTo','Z':'close'}[command[0]]
                segment=OxmlElement('a:'+tag)
                for xx,yy in zip(command[1::2],command[2::2]):
                    point=OxmlElement('a:pt'); point.set('x',str(round((xx-x)*1000))); point.set('y',str(round((yy-y)*1000))); segment.append(point)
                vector.append(segment)
            paths.append(vector); geometry.append(paths); shape._element.spPr.insert(1,geometry)
            style_shape(shape,node,scale)
        elif node['type'] == 'polygon':
            points = node['points']
            builder = slide.shapes.build_freeform(points[0][0], points[0][1], scale=Inches(scale))
            builder.add_line_segments(points[1:], close=True)
            shape = builder.convert_to_shape()
            style_shape(shape, node, scale)
        else:
            kind_value = {'rect':MSO_SHAPE.RECTANGLE,'roundrect':MSO_SHAPE.ROUNDED_RECTANGLE,'ellipse':MSO_SHAPE.OVAL}[node['type']]
            shape = slide.shapes.add_shape(kind_value, Inches(x*scale), Inches(y*scale), Inches(w*scale), Inches(h*scale))
            if node['type']=='roundrect':
                shape.adjustments[0] = min(.5,node.get('radius',12)/min(w,h))
            style_shape(shape,node,scale)
        shape.name = node['id']
        created[node['id']] = shape
        if node.get('group'): groups.setdefault(node['group'],[]).append(shape)
    for name,shapes in groups.items():
        group=slide.shapes.add_group_shape(shapes); group.name=name
    slide.notes_slide.notes_text_frame.text = scene.get('notes', 'Native shapes and editable text exported from the accompanying scene JSON.')
    prs.save(destination)


def png_from_svg(svg, destination, scene, scale):
    from playwright.sync_api import sync_playwright
    width, height = scene['canvas']['width'], scene['canvas']['height']
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={'width': math.ceil(width), 'height': math.ceil(height)}, device_scale_factor=scale)
        page.route('**/*', lambda route: route.abort())
        page.set_content('<!doctype html><meta charset="utf-8"><style>html,body{margin:0;background:white}svg{display:block}</style>' + svg)
        page.evaluate('document.fonts.ready')
        page.locator('svg').screenshot(path=str(destination))
        browser.close()


def audit_files(scene, output, warnings):
    with zipfile.ZipFile(output.with_suffix('.pptx')) as archive:
        for name in archive.namelist():
            if name.endswith(('.xml','.rels')):
                ET.fromstring(archive.read(name))
        root = ET.fromstring(archive.read('ppt/slides/slide1.xml'))
        pictures = len(root.findall('.//p:pic', NS))
        texts = root.findall('.//a:t', NS)
        expected = sum(len(node['_lines']) for node in scene['nodes'] if node['type']=='text')
        if pictures or len(texts) != expected:
            raise ValueError('PPTX editability audit failed: pictures or missing text')
        ppt_text = [item.text or '' for item in texts]
        svg_root = ET.parse(output.with_suffix('.svg')).getroot()
        svg_text = [item.text or '' for item in svg_root.findall('.//{http://www.w3.org/2000/svg}tspan')]
        if ppt_text != svg_text:
            raise ValueError('SVG and PPTX labels differ')
        native_shapes = len(root.findall('.//p:sp', NS))
        connectors = len(root.findall('.//p:cxnSp', NS))
    return {'editable_native_shapes': native_shapes, 'editable_text_runs': len(texts), 'native_connector_segments': connectors, 'raster_pictures': pictures, 'svg_pptx_labels_identical': True, 'warnings': warnings, 'node_types': dict(Counter(node['type'] for node in scene['nodes'])), 'limitations': ['PPTX text and geometry are native; visual rendering in desktop Office has not been checked by this script.', 'Check overlap, arrow routing, semantics, and final-size readability visually.']}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('scene', type=Path)
    parser.add_argument('--output', required=True, type=Path, help='Output basename, e.g. output/figure')
    parser.add_argument('--png-scale', type=float, default=2)
    parser.add_argument('--force', action='store_true', help='Explicitly replace existing outputs')
    args = parser.parse_args()
    if not 0 < args.png_scale <= 8:
        parser.error('--png-scale must be >0 and <=8')
    output = args.output.resolve()
    for suffix in ('.svg','.pptx','.png','.qa.json'):
        if output.with_suffix(suffix).exists() and not args.force:
            parser.error(f'{output.with_suffix(suffix)} exists; choose another name or --force')
    scene, warnings = prepare(json.loads(args.scene.read_text(encoding='utf-8-sig')))
    output.parent.mkdir(parents=True,exist_ok=True)
    svg = build_svg(scene)
    output.with_suffix('.svg').write_text(svg,encoding='utf-8')
    build_pptx(scene,output.with_suffix('.pptx'))
    png_from_svg(svg,output.with_suffix('.png'),scene,args.png_scale)
    report = audit_files(scene,output,warnings)
    output.with_suffix('.qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output': str(output),'qa':report},ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
